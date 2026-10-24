"""Small, deterministic source adapters; PDF extraction produces review candidates only."""

import calendar
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any, Literal
from urllib.parse import parse_qs, urlsplit

from pydantic import BaseModel, ConfigDict, Field, model_validator

from shiftos.services.evidence.registry import RegistryEntry

PARSER_VERSION = "phase3-evidence-1"


class EvidenceParseError(ValueError):
    """A visible layout/structure failure, never an implicit empty successful import."""


class Fact(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    metadata_json: dict[str, Any] = Field(default_factory=dict)


class OccupationData(Fact):
    taxonomy: str = Field(min_length=1)
    label: str = Field(min_length=1)
    description: str | None = None
    classification_code: str | None = None


class SkillData(Fact):
    taxonomy: str = Field(min_length=1)
    label: str = Field(min_length=1)
    description: str | None = None


class RelationData(Fact):
    occupation_external_id: str = Field(min_length=1)
    skill_external_id: str = Field(min_length=1)
    relationship: Literal["essential", "optional"]


class TaskData(Fact):
    occupation_external_id: str | None = None
    description: str = Field(min_length=1)


class WorkforceData(Fact):
    occupation_external_id: str | None = None
    occupation_label: str | None = None
    classification_system: str | None = None
    classification_code: str | None = None
    sector: str | None = None
    geography: str = Field(min_length=1)
    period_start: date
    period_end: date
    indicator: str = Field(min_length=1)
    value: Decimal | None
    unit: str = Field(min_length=1)
    value_status: str = "published"
    limitations: str | None = None

    @model_validator(mode="after")
    def coherent_period(self):
        if self.period_end < self.period_start:
            raise ValueError("Observation period end precedes start")
        if self.value is not None and not self.value.is_finite():
            raise ValueError("Workforce value must be finite")
        if not self.occupation_label and not self.sector and not self.occupation_external_id:
            raise ValueError("Workforce observation requires a published population")
        return self


class AIData(Fact):
    occupation_label: str | None = None
    task: str | None = None
    statement: str = Field(min_length=1)
    impact_type: Literal[
        "observed_capability", "observed_practice", "occupational_exposure", "future_prediction"
    ]
    capability: str | None = None
    limitations: str = Field(min_length=1)
    study_period: str | None = None


class AcademicData(Fact):
    discipline: str = Field(min_length=1)
    title: str = Field(min_length=1)
    statement: str = Field(min_length=1)
    standard_code: str | None = None


class SDGData(Fact):
    code: str = Field(pattern=r"^\d{1,2}(?:\.(?:\d{1,2}|[a-z]))?$")
    kind: Literal["goal", "target"]
    goal_code: str
    title: str = Field(min_length=1)
    description: str | None = None

    @model_validator(mode="after")
    def coherent_goal(self):
        if not self.goal_code.isdigit() or not 1 <= int(self.goal_code) <= 17:
            raise ValueError("SDG goal must be in official range 1..17")
        if self.code.split(".")[0] != self.goal_code:
            raise ValueError("SDG target belongs to a different goal")
        if ("." in self.code) != (self.kind == "target"):
            raise ValueError("SDG code/kind mismatch")
        return self


class PDFData(Fact):
    text: str
    tables: list[list[list[str | None]]] = Field(default_factory=list)


DATA_SCHEMAS = {
    "occupation": OccupationData,
    "skill": SkillData,
    "occupation_skill": RelationData,
    "task": TaskData,
    "workforce": WorkforceData,
    "ai_evidence": AIData,
    "academic_standard": AcademicData,
    "sdg": SDGData,
    "pdf_candidate": PDFData,
}


class ParsedRecord(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    kind: Literal[
        "occupation",
        "skill",
        "occupation_skill",
        "task",
        "workforce",
        "ai_evidence",
        "academic_standard",
        "sdg",
        "pdf_candidate",
    ]
    key: str = Field(min_length=1, max_length=1000)
    data: dict[str, Any]
    locator: str = Field(min_length=1, max_length=1000)
    excerpt: str = Field(min_length=1)
    page_number: int | None = Field(default=None, ge=1)
    review_status: Literal["source_assertion", "reviewed", "needs_review"] = "source_assertion"

    @model_validator(mode="after")
    def validate_typed_data(self):
        self.data = DATA_SCHEMAS[self.kind].model_validate(self.data).model_dump(mode="json")
        if self.kind == "pdf_candidate" and self.review_status != "needs_review":
            raise ValueError("Automatically extracted PDF text must remain needs_review")
        return self


class ParsedBatch(BaseModel):
    records: list[ParsedRecord]
    metadata: dict[str, Any] = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)

    @model_validator(mode="after")
    def safe_duplicates(self):
        unique: dict[tuple[str, str], ParsedRecord] = {}
        for record in self.records:
            identity = (record.kind, record.key)
            previous = unique.get(identity)
            if previous and previous.data != record.data:
                raise ValueError(f"Conflicting duplicate record: {identity}")
            unique.setdefault(identity, record)
        removed = len(self.records) - len(unique)
        if removed:
            self.warnings.append(f"Removed {removed} exact duplicate source records")
        self.records = list(unique.values())
        self.metadata.setdefault("parser_version", PARSER_VERSION)
        return self


def _json(content: bytes) -> Any:
    try:
        return json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise EvidenceParseError(
            "Invalid UTF-8 JSON; HTML landing/error pages are not data"
        ) from exc


def _csv(content: bytes, required: set[str]) -> list[dict[str, str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise EvidenceParseError("CSV must be UTF-8") from exc
    reader = csv.DictReader(io.StringIO(text))
    if not required <= set(reader.fieldnames or []):
        raise EvidenceParseError(f"CSV missing required columns: {sorted(required)}")
    rows = list(reader)
    if not rows or any(None in row or any(v is None for v in row.values()) for row in rows):
        raise EvidenceParseError("CSV is empty or has inconsistent columns")
    return rows


def _number(value: Any) -> tuple[Decimal | None, str]:
    if value is None or str(value).strip() in {"", "-", "..", "...", "NA", "N/A"}:
        return None, "missing_published_sentinel"
    try:
        number = Decimal(str(value).replace(",", "").strip())
        if not number.is_finite():
            raise InvalidOperation
        return number, "published"
    except InvalidOperation as exc:
        raise EvidenceParseError(f"Invalid published number: {value!r}") from exc


def parse_employment_sector(content: bytes) -> ParsedBatch:
    records = []
    for number, row in enumerate(_csv(content, {"date", "sector", "sex", "proportion"}), 2):
        try:
            encoded = date.fromisoformat(row["date"])
        except ValueError as exc:
            raise EvidenceParseError(f"Invalid annual period at CSV row {number}") from exc
        if encoded.month != 1 or encoded.day != 1:
            raise EvidenceParseError("OpenDOSM annual period encoding changed; review required")
        value, status = _number(row["proportion"])
        if value is not None and not 0 <= value <= 100:
            raise EvidenceParseError("Employment proportion is outside 0..100 percent")
        records.append(
            ParsedRecord(
                kind="workforce",
                key=f"{row['date']}:{row['sector']}:{row['sex']}",
                locator=f"CSV row {number}",
                excerpt=", ".join(f"{k}={v}" for k, v in row.items()),
                data={
                    "sector": row["sector"],
                    "geography": "MY",
                    "period_start": encoded,
                    "period_end": date(encoded.year, 12, 31),
                    "indicator": "employment_distribution",
                    "value": value,
                    "unit": "percent",
                    "value_status": status,
                    "limitations": (
                        "Broad survey sector distribution; not job vacancies or skill demand."
                    ),
                    "metadata_json": {
                        "sex": row["sex"],
                        "raw": row,
                        "period_encoding": "publisher January 1 denotes annual period",
                    },
                },
            )
        )
    return ParsedBatch(records=records)


def _workbook(content: bytes):
    from openpyxl import load_workbook

    if not content.startswith(b"PK"):
        raise EvidenceParseError("Expected XLSX ZIP signature; download may be a landing page")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 5000
                or sum(item.file_size for item in members) > 100_000_000
                or any(item.file_size > 30_000_000 for item in members)
            ):
                raise EvidenceParseError("XLSX exceeds bounded uncompressed archive limits")
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except EvidenceParseError:
        raise
    except Exception as exc:
        raise EvidenceParseError("Corrupted or unsupported XLSX workbook") from exc
    if (
        any((sheet.max_row or 0) > 5000 or (sheet.max_column or 0) > 500 for sheet in workbook)
        or sum((sheet.max_row or 0) * (sheet.max_column or 0) for sheet in workbook) > 2_000_000
    ):
        workbook.close()
        raise EvidenceParseError("Workbook exceeds bounded cell limit")
    return workbook


def parse_dosm_jmi(content: bytes, dataset: RegistryEntry | None = None) -> ParsedBatch:
    """Find the occupation table and merged year/quarter headers, never YTD columns."""
    workbook = _workbook(content)
    records = []
    try:
        sheet = next(
            (s for s in workbook if "summary" in s.title.lower() and "bi" in s.title.lower()), None
        )
        if sheet is None:
            raise EvidenceParseError("DOSM English summary sheet not found; layout changed")
        tables = [
            (cell.row, cell.column)
            for row in sheet
            for cell in row
            if isinstance(cell.value, str)
            and re.search(r"Table\s*2\s*:.*Job Offerings.*Occupation", cell.value, re.I)
        ]
        if len(tables) != 1:
            raise EvidenceParseError("Expected one DOSM Table 2 occupation count table")
        title_row, label_column = tables[0]
        header_rows = [
            r
            for r in range(title_row + 1, min(title_row + 8, sheet.max_row) + 1)
            if str(sheet.cell(r, label_column).value).strip().lower() == "occupation"
        ]
        if len(header_rows) != 1:
            raise EvidenceParseError("DOSM occupation/quarter header not found")
        quarter_row = header_rows[0]
        selected = re.search(r"Q([1-4])-(\d{4})", dataset.id) if dataset else None
        periods = []
        year = None
        for col in range(label_column + 1, sheet.max_column + 1):
            raw_year = sheet.cell(quarter_row - 1, col).value
            if isinstance(raw_year, (int, float)) and 1900 <= raw_year <= 2100:
                year = int(raw_year)
            quarter = str(sheet.cell(quarter_row, col).value).strip()
            if re.fullmatch(r"Q[1-4]", quarter) and year is not None:
                q = int(quarter[1])
                if selected and (q != int(selected[1]) or year != int(selected[2])):
                    continue
                periods.append((col, year, q))
        if not periods:
            raise EvidenceParseError("Requested DOSM observation quarter not present")
        for row in range(quarter_row + 1, sheet.max_row + 1):
            raw_label = sheet.cell(row, label_column).value
            if raw_label is None or str(raw_label).startswith("Source:"):
                break
            label = str(raw_label).strip()
            for col, year, quarter in periods:
                raw_value = sheet.cell(row, col).value
                # Empty future quarter cells are not observations; explicit dash/zero retained.
                if raw_value is None:
                    continue
                value, status = _number(raw_value)
                start_month = 3 * quarter - 2
                end_month = start_month + 2
                cell = sheet.cell(row, col).coordinate
                records.append(
                    ParsedRecord(
                        kind="workforce",
                        key=f"{year}-Q{quarter}:{label}",
                        locator=(
                            f"{sheet.title}!{cell}; label "
                            f"{sheet.cell(row, label_column).coordinate}; Table 2"
                        ),
                        excerpt=f"{label}; {year} Q{quarter}; {raw_value} online job offerings",
                        data={
                            "occupation_label": label,
                            "geography": "MY",
                            "period_start": date(year, start_month, 1),
                            "period_end": date(
                                year, end_month, calendar.monthrange(year, end_month)[1]
                            ),
                            "indicator": "online_job_offerings",
                            "value": value,
                            "unit": "job_offerings",
                            "value_status": status,
                            "limitations": (
                                "Major occupation groups/total; online advertising scope only. "
                                "Not unique hires, individual occupations or skill demand; "
                                "no MASCO/ISCO crosswalk has been inferred."
                            ),
                            "metadata_json": {
                                "raw_value": raw_value,
                                "raw_label": label,
                                "table_title": sheet.cell(title_row, label_column).value,
                                "classification_status": "unmatched",
                            },
                        },
                    )
                )
        if not records:
            raise EvidenceParseError("DOSM table has no published values in requested period")
        return ParsedBatch(records=records, metadata={"sheet": sheet.title, "table": "Table 2"})
    finally:
        workbook.close()


def parse_un_sdg(content: bytes) -> ParsedBatch:
    payload = _json(content)
    if not isinstance(payload, list) or not payload:
        raise EvidenceParseError("Expected nonempty official UN goal/target JSON list")
    records = []
    for row in payload:
        if not isinstance(row, dict) or not {"code", "title", "description", "uri"} <= row.keys():
            raise EvidenceParseError("UN goal/target JSON schema changed")
        code = str(row["code"])
        kind = "target" if "goal" in row else "goal"
        records.append(
            ParsedRecord(
                kind="sdg",
                key=f"{kind}:{code}",
                locator=str(row["uri"]),
                excerpt=row["title"],
                data={
                    "code": code,
                    "kind": kind,
                    "goal_code": str(row.get("goal", code)),
                    "title": row["title"],
                    "description": row["description"],
                    "metadata_json": {"publisher_uri": row["uri"]},
                },
            )
        )
    return ParsedBatch(
        records=records,
        warnings=[
            "Official reference semantics only; no occupation/skill-to-SDG relationship inferred."
        ],
    )


def _english(value: Any) -> str | None:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        selected = value.get("en", value.get("literal"))
        return _english(selected) if selected is not None else None
    return None


def _concept_uri(link: dict[str, Any]) -> str | None:
    if link.get("uri"):
        return link["uri"]
    query = parse_qs(urlsplit(str(link.get("href", ""))).query)
    return query.get("uri", [None])[0]


def parse_esco_api(content: bytes) -> ParsedBatch:
    payload = _json(content)
    if not isinstance(payload, dict):
        raise EvidenceParseError("ESCO API must return a JSON object")
    concepts = payload.get("_embedded", {}).get("results")
    if concepts is None:
        concepts = [payload]
    records = []
    for row in concepts:
        if not isinstance(row, dict):
            raise EvidenceParseError("Invalid ESCO concept")
        uri = row.get("uri")
        label = _english(row.get("preferredLabel")) or _english(row.get("title"))
        kind = row.get("className", row.get("type"))
        if kind not in {"Occupation", "Skill", "occupation", "skill"} or not uri or not label:
            raise EvidenceParseError("ESCO concept requires actual URI, label and explicit type")
        kind = kind.lower()
        data = {
            "taxonomy": "ESCO",
            "label": label,
            "description": _english(row.get("description")),
            "metadata_json": {"raw_type": row.get("className", row.get("type"))},
        }
        if kind == "occupation":
            data["classification_code"] = row.get("code")
        records.append(ParsedRecord(kind=kind, key=uri, data=data, locator=uri, excerpt=label))
        if kind == "occupation":
            for link_name, relationship in [
                ("hasEssentialSkill", "essential"),
                ("hasOptionalSkill", "optional"),
            ]:
                links = row.get("_links", {}).get(link_name, [])
                if isinstance(links, dict):
                    links = [links]
                for index, link in enumerate(links):
                    skill_uri, skill_label = _concept_uri(link), _english(link.get("title"))
                    if not skill_uri or not skill_label:
                        raise EvidenceParseError(
                            "ESCO skill relation lacks source URI/title; fetch full resource"
                        )
                    records.append(
                        ParsedRecord(
                            kind="skill",
                            key=skill_uri,
                            locator=f"{uri}#/_links/{link_name}/{index}/title",
                            excerpt=skill_label,
                            data={"taxonomy": "ESCO", "label": skill_label},
                        )
                    )
                    records.append(
                        ParsedRecord(
                            kind="occupation_skill",
                            key=f"{uri}:{relationship}:{skill_uri}",
                            locator=f"{uri}#/_links/{link_name}/{index}",
                            excerpt=f"{label}: {relationship} skill {skill_label}",
                            data={
                                "occupation_external_id": uri,
                                "skill_external_id": skill_uri,
                                "relationship": relationship,
                            },
                        )
                    )
    if not records:
        raise EvidenceParseError("ESCO API returned no concepts")
    return ParsedBatch(
        records=records,
        warnings=[
            "ESCO is an EU occupational benchmark, "
            "not Malaysian hiring or university teaching evidence."
        ],
    )


def parse_esco_csv(content: bytes, adapter: str) -> ParsedBatch:
    relation = adapter == "esco_relations_csv"
    required = (
        {"occupationUri", "skillUri", "relationType"}
        if relation
        else {"conceptUri", "preferredLabel"}
    )
    records = []
    for number, row in enumerate(_csv(content, required), 2):
        if relation:
            occupation, skill, relationship = (
                row["occupationUri"],
                row["skillUri"],
                row["relationType"],
            )
            records.append(
                ParsedRecord(
                    kind="occupation_skill",
                    key=f"{occupation}:{relationship}:{skill}",
                    locator=f"CSV row {number}",
                    excerpt=f"{occupation} {relationship} {skill}",
                    data={
                        "occupation_external_id": occupation,
                        "skill_external_id": skill,
                        "relationship": relationship,
                        "metadata_json": {"raw": row},
                    },
                )
            )
        else:
            kind = "occupation" if adapter == "esco_occupations_csv" else "skill"
            data = {
                "taxonomy": "ESCO",
                "label": row["preferredLabel"],
                "description": row.get("description"),
                "metadata_json": {"raw": row},
            }
            if kind == "occupation":
                data["classification_code"] = row.get("code")
            records.append(
                ParsedRecord(
                    kind=kind,
                    key=row["conceptUri"],
                    data=data,
                    locator=f"CSV row {number}",
                    excerpt=row["preferredLabel"],
                )
            )
    return ParsedBatch(records=records)


def parse_reviewed_records(content: bytes) -> ParsedBatch:
    payload = _json(content)
    if not isinstance(payload, dict) or payload.get("reviewed") is not True:
        raise EvidenceParseError("Manual fact import requires an explicit reviewed=true assertion")
    if not isinstance(payload.get("records"), list) or not payload["records"]:
        raise EvidenceParseError("Reviewed import requires nonempty records")
    batch = ParsedBatch.model_validate(
        {"records": payload["records"], "metadata": payload.get("metadata", {})}
    )
    if any(record.review_status != "reviewed" for record in batch.records):
        raise EvidenceParseError("Each manual fact must carry review_status=reviewed")
    batch.warnings.append(
        "Reviewed is an operator assertion; "
        "it does not independently verify factual interpretation."
    )
    return batch


def parse_pdf_candidates(content: bytes, dataset: RegistryEntry | None = None) -> ParsedBatch:
    import pdfplumber

    if not content.startswith(b"%PDF-"):
        raise EvidenceParseError("Expected PDF signature")
    records, warnings = [], []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages[:120]:
                text = page.extract_text(layout=False) or ""
                if not text.strip():
                    warnings.append(
                        f"Page {page.page_number}: no text layer; manual/OCR review required"
                    )
                    continue
                tables = page.extract_tables() or []
                records.append(
                    ParsedRecord(
                        kind="pdf_candidate",
                        key=f"page:{page.page_number}",
                        locator=f"PDF page {page.page_number}",
                        page_number=page.page_number,
                        review_status="needs_review",
                        excerpt=text[:2000],
                        data={
                            "text": text[:100000],
                            "tables": [table[:200] for table in tables[:10]],
                        },
                    )
                )
            if len(pdf.pages) > 120:
                warnings.append(
                    "Document exceeds 120-page extraction bound; "
                    "remaining pages require targeted manual extraction"
                )
            metadata = {
                "pdf_metadata": {k: str(v) for k, v in (pdf.metadata or {}).items()},
                "total_pages": len(pdf.pages),
                "document_title": dataset.name if dataset else None,
                "publisher": dataset.publisher if dataset else None,
                "publication_date_raw": dataset.latest_verified_publication_date
                if dataset
                else None,
            }
    except Exception as exc:
        raise EvidenceParseError("PDF extraction failed; manual extraction required") from exc
    if not records:
        raise EvidenceParseError("PDF has no extractable text; manual/OCR review required")
    warnings.append(
        "PDF text/table extraction candidates are not verified factual findings "
        "or occupation links."
    )
    return ParsedBatch(records=records, metadata=metadata, warnings=warnings)


ADAPTERS = {
    "dosm_jmi",
    "employment_sector",
    "esco_api",
    "esco_occupations_csv",
    "esco_skills_csv",
    "esco_relations_csv",
    "un_sdg",
    "reviewed_json",
    "pdf_candidates",
}


def parse_dataset(
    content: bytes, *, format: str, adapter: str, dataset: RegistryEntry | None = None
) -> ParsedBatch:
    if not content:
        raise EvidenceParseError("Empty source file")
    if len(content) > 30_000_000:
        raise EvidenceParseError("Source file exceeds 30MB parsing bound")
    expected_formats = {
        "dosm_jmi": "xlsx",
        "employment_sector": "csv",
        "esco_api": "json",
        "un_sdg": "json",
        "reviewed_json": "json",
        "pdf_candidates": "pdf",
    }
    expected = (
        "csv"
        if adapter.startswith("esco_") and adapter.endswith("_csv")
        else expected_formats.get(adapter)
    )
    if adapter not in ADAPTERS or expected != format.lower():
        raise EvidenceParseError(
            "Unsupported adapter/format combination; manual transformation required"
        )
    if adapter == "dosm_jmi":
        result = parse_dosm_jmi(content, dataset)
    elif adapter == "employment_sector":
        result = parse_employment_sector(content)
    elif adapter == "un_sdg":
        result = parse_un_sdg(content)
    elif adapter == "esco_api":
        result = parse_esco_api(content)
    elif adapter.endswith("_csv"):
        result = parse_esco_csv(content, adapter)
    elif adapter == "pdf_candidates":
        result = parse_pdf_candidates(content, dataset)
    else:
        result = parse_reviewed_records(content)
    result.metadata.update({"adapter": adapter, "sha256": hashlib.sha256(content).hexdigest()})
    return result
