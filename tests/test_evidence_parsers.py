"""All in-memory source contents below are explicitly synthetic test fixtures."""

import csv
import io
import json
import struct
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from shiftos.services.evidence.parsers import (
    EvidenceParseError,
    ParsedBatch,
    ParsedRecord,
    parse_dataset,
    parse_dosm_jmi,
    parse_employment_sector,
    parse_esco_api,
    parse_esco_csv,
    parse_pdf_candidates,
    parse_reviewed_records,
    parse_un_sdg,
)
from shiftos.services.evidence.registry import REGISTRY_COLUMNS, read_registry

ROOT = Path(__file__).resolve().parents[1]


def _entry(dataset_id="D-JMI-Q1-2026"):
    return next(
        row for row in read_registry(ROOT / "data/dataset_registry.csv") if row.id == dataset_id
    )


def _write_registry(tmp_path, rows, columns=None):
    path = tmp_path / "registry.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        fieldnames = columns or list(
            dict.fromkeys([*REGISTRY_COLUMNS, *(key for row in rows for key in row)])
        )
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return path


def test_real_registry_retains_research_core_and_accepts_verified_additions():
    rows = read_registry(ROOT / "data/dataset_registry.csv")
    assert len(rows) == len({row.id for row in rows}) >= 73
    assert {"C-ESCO", "F-UN-SDG", "D-JMI-Q1-2026", "B-MQA-COMPUTING"} <= {r.id for r in rows}
    assert _entry("B-MQA-DIRECTORY").download_urls == []
    assert _entry().publication_date.isoformat() == "2026-06-25"
    assert _entry("F-UN-SDG").publication_date is None  # Year-only date not invented Jan 1.


def test_registry_rejects_missing_columns(tmp_path):
    with pytest.raises(ValueError, match="missing columns"):
        read_registry(_write_registry(tmp_path, [{"id": "X"}], columns=["id"]))


def test_registry_rejects_duplicate_ids(tmp_path):
    row = _entry().model_dump(mode="json")
    with pytest.raises(ValueError, match="duplicate dataset ID"):
        read_registry(_write_registry(tmp_path, [row, row]))


@pytest.mark.parametrize(
    "field,value",
    [
        ("source_url", "https://user:secret@example.invalid/data"),
        ("source_url", "not-a-url"),
        ("download_api_url", "ftp://example.invalid/test"),
        ("evidence_class", "proved_skill_demand"),
        ("priority", "P9"),
        ("verification_status", "Acquired"),
    ],
)
def test_registry_rejects_invalid_metadata(tmp_path, field, value):
    row = _entry().model_dump(mode="json")
    row[field] = value
    with pytest.raises(ValidationError):
        read_registry(_write_registry(tmp_path, [row]))


def _xlsx(shift=0, quarter="Q1", raw_value=123, label="Synthetic science professionals"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "e2-Summary Table BI"
    row = 3 + shift
    sheet.cell(
        row, 2, "Table 2: Number of Job Offerings Advertised Online by Major Group Occupation"
    )
    sheet.cell(row + 2, 3, 2026)
    sheet.cell(row + 3, 2, "Occupation")
    sheet.cell(row + 3, 3, quarter)
    sheet.cell(row + 3, 4, "Q4")  # Merged year cell extends horizontally.
    sheet.cell(row + 4, 2, label)
    sheet.cell(row + 4, 3, raw_value)
    sheet.cell(row + 4, 4, None)
    sheet.cell(row + 5, 2, "Source: Synthetic fixture")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_dosm_layout_shift_year_period_unit_and_original_value():
    batch = parse_dosm_jmi(_xlsx(shift=7), _entry())
    assert len(batch.records) == 1
    record = batch.records[0]
    assert record.data["period_start"] == "2026-01-01"
    assert record.data["period_end"] == "2026-03-31"
    assert record.data["unit"] == "job_offerings"
    assert record.data["value"] == "123"
    assert record.data["classification_code"] is None
    assert record.data["metadata_json"]["raw_value"] == 123
    assert record.locator.startswith("e2-Summary Table BI!C14")


@pytest.mark.parametrize(
    "raw_value,expected,status",
    [
        ("-", None, "missing_published_sentinel"),
        (0, "0", "published"),
        ("1,234", "1234", "published"),
    ],
)
def test_dosm_dash_and_zero_are_not_conflated(raw_value, expected, status):
    item = parse_dosm_jmi(_xlsx(raw_value=raw_value), _entry()).records[0]
    assert item.data["value"] == expected
    assert item.data["value_status"] == status


def test_dosm_q2_has_separate_observation_window():
    item = parse_dosm_jmi(_xlsx(quarter="Q2"), _entry("D-JMI-Q2-2026")).records[0]
    assert item.data["period_start"] == "2026-04-01"
    assert item.data["period_end"] == "2026-06-30"


def test_dosm_fails_if_requested_quarter_absent():
    with pytest.raises(EvidenceParseError, match="quarter not present"):
        parse_dosm_jmi(_xlsx(quarter="Q2"), _entry())


@pytest.mark.parametrize("content", [b"<html>Landing page</html>", b"PKbroken"])
def test_invalid_excel_integrity_is_visible(content):
    with pytest.raises(EvidenceParseError):
        parse_dosm_jmi(content)


def test_xlsx_expanded_zip_size_is_checked_before_openpyxl():
    # Synthetic archive metadata advertises an oversized entry without allocating it.
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("synthetic.xml", b"x")
    content = bytearray(stream.getvalue())
    directory = content.index(b"PK\x01\x02")
    struct.pack_into("<I", content, directory + 24, 30_000_001)
    with pytest.raises(EvidenceParseError, match="uncompressed archive limits"):
        parse_dosm_jmi(bytes(content))


def test_employment_sector_preserves_historical_rows_units_and_sex():
    batch = parse_employment_sector(
        b"date,sector,sex,proportion\n2024-01-01,synthetic-sector,both,10.5\n"
        b"2025-01-01,synthetic-sector,both,11.2\n"
    )
    assert len(batch.records) == 2
    assert batch.records[0].data["period_end"] == "2024-12-31"
    assert batch.records[0].data["unit"] == "percent"
    assert batch.records[0].data["metadata_json"]["sex"] == "both"
    assert batch.records[0].data["occupation_external_id"] is None


@pytest.mark.parametrize(
    "content",
    [
        b"date,sector,sex,proportion\n2024-03-01,x,both,10\n",
        b"date,sector,sex,proportion\n2024-01-01,x,both,101\n",
        b"date,sector,sex,proportion\n2024-01-01,x,both,NaN\n",
        b"date,sector,sex,proportion\n2024-01-01,x\n",
        b"unexpected_header\nvalue\n",
        b"\xff",
        b"date,sector,sex,proportion\n",
    ],
)
def test_corrupted_or_changed_csv_fails(content):
    with pytest.raises((EvidenceParseError, ValidationError)):
        parse_employment_sector(content)


def test_un_goals_and_targets_preserve_official_ids_without_mapping():
    row = {
        "code": "7.a",
        "goal": "7",
        "title": "Synthetic UN target fixture",
        "description": "Explicitly synthetic",
        "uri": "/v1/sdg/Target/7.a",
        "indicators": None,
    }
    record = parse_un_sdg(json.dumps([row]).encode()).records[0]
    assert record.key == "target:7.a"
    assert record.data["goal_code"] == "7"
    assert not any("skill" in key for key in record.data)


def test_un_goal_code_consistency_is_validated():
    row = {
        "code": "7.2",
        "goal": "8",
        "title": "Synthetic mismatch",
        "description": "Synthetic",
        "uri": "/synthetic",
    }
    with pytest.raises(ValidationError, match="different goal"):
        parse_un_sdg(json.dumps([row]).encode())


def test_esco_real_uri_links_required_not_guessed_by_label():
    payload = {
        "uri": "http://example.invalid/occupation/synthetic-accountant",
        "className": "Occupation",
        "preferredLabel": {"en": "Synthetic accountant"},
        "description": {"en": {"literal": "Synthetic occupation description"}},
        "code": "SYNTH",
        "_links": {
            "hasEssentialSkill": [
                {"uri": "http://example.invalid/skill/synthetic-audit", "title": "Synthetic audit"}
            ]
        },
    }
    batch = parse_esco_api(json.dumps(payload).encode())
    assert [item.kind for item in batch.records] == ["occupation", "skill", "occupation_skill"]
    assert batch.records[-1].data["relationship"] == "essential"
    assert batch.records[0].data["description"] == "Synthetic occupation description"
    payload["_links"]["hasEssentialSkill"][0].pop("uri")
    with pytest.raises(EvidenceParseError, match="lacks source URI/title"):
        parse_esco_api(json.dumps(payload).encode())


def test_esco_search_results_and_no_fabricated_skill_relations():
    payload = {
        "_embedded": {
            "results": [
                {
                    "uri": "http://example.invalid/occupation/scientist",
                    "className": "Occupation",
                    "title": "Synthetic scientist",
                },
                {
                    "uri": "http://example.invalid/occupation/editor",
                    "className": "Occupation",
                    "title": "Synthetic editor",
                },
            ]
        }
    }
    batch = parse_esco_api(json.dumps(payload).encode())
    assert len(batch.records) == 2
    assert all(item.kind == "occupation" for item in batch.records)


def test_esco_manual_csv_stable_id_join_preserves_source_fields():
    batch = parse_esco_csv(
        b"conceptUri,preferredLabel,description,code\n"
        b"http://example.invalid/occupation/synthetic,Synthetic scientist,Test,SYN\n",
        "esco_occupations_csv",
    )
    assert batch.records[0].data["classification_code"] == "SYN"
    assert batch.records[0].data["metadata_json"]["raw"]["description"] == "Test"
    batch = parse_esco_csv(
        b"occupationUri,skillUri,relationType\n"
        b"http://example.invalid/occupation/synthetic,http://example.invalid/skill/x,optional\n",
        "esco_relations_csv",
    )
    assert batch.records[0].data["relationship"] == "optional"


def _manual_record():
    return {
        "kind": "ai_evidence",
        "key": "synthetic-impact-1",
        "locator": "PDF p. 4, Table 1",
        "page_number": 4,
        "excerpt": "Synthetic exposure model test excerpt",
        "review_status": "reviewed",
        "data": {
            "statement": "Synthetic exposure fixture, not actual research",
            "impact_type": "occupational_exposure",
            "limitations": "Explicitly synthetic",
            "task": "Synthetic task",
            "study_period": "2025",
        },
    }


def test_reviewed_cards_require_operator_review_excerpts_and_ai_limitations():
    record = _manual_record()
    payload = {"reviewed": True, "records": [record]}
    assert parse_reviewed_records(json.dumps(payload).encode()).records[0].page_number == 4
    record["data"].pop("limitations")
    with pytest.raises(ValidationError):
        parse_reviewed_records(json.dumps(payload).encode())


@pytest.mark.parametrize(
    "field,value", [("excerpt", ""), ("locator", ""), ("review_status", "needs_review")]
)
def test_manual_cards_cannot_omit_provenance_or_review(field, value):
    record = _manual_record()
    record[field] = value
    with pytest.raises((EvidenceParseError, ValidationError)):
        parse_reviewed_records(json.dumps({"reviewed": True, "records": [record]}).encode())


def test_future_predictions_cannot_be_mislabelled_automatic_replacement_probability():
    record = _manual_record()
    record["data"]["impact_type"] = "replacement_probability"
    with pytest.raises(ValidationError):
        ParsedRecord.model_validate(record)


def test_equal_duplicates_deduplicate_conflicting_duplicate_fails():
    item = ParsedRecord.model_validate(_manual_record())
    batch = ParsedBatch(records=[item, item.model_copy(deep=True)])
    assert len(batch.records) == 1
    assert "duplicate" in batch.warnings[0]
    conflict = item.model_copy(deep=True)
    conflict.data["statement"] = "Different synthetic statement"
    with pytest.raises(ValidationError, match="Conflicting duplicate"):
        ParsedBatch(records=[item, conflict])


def _synthetic_pdf():
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = b"BT /F1 12 Tf 72 720 Td (Synthetic research fixture only) Tj ET"
    objects.append(
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    )
    output, offsets = b"%PDF-1.4\n", [0]
    for i, value in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{i} 0 obj\n".encode() + value + b"\nendobj\n"
    start = len(output)
    output += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    output += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    output += f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    return output


def test_pdf_extraction_retains_page_metadata_but_only_review_candidates():
    batch = parse_pdf_candidates(_synthetic_pdf(), _entry("E-ILO-2025"))
    assert len(batch.records) == 1
    record = batch.records[0]
    assert record.kind == "pdf_candidate" and record.review_status == "needs_review"
    assert record.page_number == 1
    assert "Synthetic research" in record.excerpt
    assert batch.metadata["publisher"] == _entry("E-ILO-2025").publisher
    assert batch.metadata["publication_date_raw"]


@pytest.mark.parametrize("content", [b"not PDF", b"%PDF-broken"])
def test_invalid_pdf_fails_visibly(content):
    with pytest.raises(EvidenceParseError):
        parse_pdf_candidates(content)


@pytest.mark.parametrize(
    "content,format,adapter",
    [
        (b"", "csv", "employment_sector"),
        (b"<html>Denied</html>", "json", "esco_api"),
        (b"[]", "parquet", "un_sdg"),
        (b"[]", "json", "unknown"),
        (b"[]", "json", "un_sdg"),
    ],
)
def test_unsupported_empty_or_error_responses_are_never_silent_success(content, format, adapter):
    with pytest.raises(EvidenceParseError):
        parse_dataset(content, format=format, adapter=adapter)


def test_parser_reports_adapter_and_checksum():
    content = b"date,sector,sex,proportion\n2024-01-01,synthetic,both,10.5\n"
    result = parse_dataset(content, format="csv", adapter="employment_sector")
    assert result.metadata["adapter"] == "employment_sector"
    assert len(result.metadata["sha256"]) == 64
